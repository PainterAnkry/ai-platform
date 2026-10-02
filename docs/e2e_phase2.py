"""Phase 2 end-to-end: register -> KB CRUD -> upload real files -> parse -> tasks.

Runs against a live backend. Exercises the inline fallback path (no Redis),
which is exactly what happens in a local preview.

Uploads all three supported formats (TXT / DOCX / PDF) through the real
multipart endpoint, so the router -> service -> parser -> chunker chain is
covered end to end, then checks the task list, the dashboard and cascade
deletion.

Usage (backend running on :8000):
    python docs/e2e_phase2.py [--pdf path/to/real.pdf]
"""
from __future__ import annotations

import argparse
import io
import json
import time
import urllib.error
import urllib.request
import uuid

BASE = "http://127.0.0.1:8000/api/v1"

parser = argparse.ArgumentParser()
parser.add_argument("--pdf", default="", help="extra real-world PDF to upload as well")
args = parser.parse_args()


# ---------------------------------------------------------------------------
# Fixtures generated in memory — no binary files are committed
# ---------------------------------------------------------------------------
def build_docx(paragraphs: list[str]) -> bytes:
    import docx

    document = docx.Document()
    document.add_heading("季度销售报告", level=1)
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_pdf(page_texts: list[str]) -> bytes:
    """Assemble a minimal valid PDF (correct xref table, Helvetica base font)."""
    objects: list[bytes] = []

    def add(body: bytes) -> int:
        objects.append(body)
        return len(objects)

    catalog_num, pages_num = add(b""), add(b"")
    font_num = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    page_nums: list[int] = []
    for text in page_texts:
        escaped = text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
        content = f"BT /F1 18 Tf 72 700 Td ({escaped}) Tj ET".encode("latin-1") if text else b""
        content_num = add(b"<< /Length %d >>\nstream\n%s\nendstream" % (len(content), content))
        page_nums.append(
            add(
                (
                    "<< /Type /Page /Parent %d 0 R /MediaBox [0 0 612 792] "
                    "/Resources << /Font << /F1 %d 0 R >> >> /Contents %d 0 R >>"
                    % (pages_num, font_num, content_num)
                ).encode("latin-1")
            )
        )

    kids = b" ".join(b"%d 0 R" % num for num in page_nums)
    objects[catalog_num - 1] = b"<< /Type /Catalog /Pages %d 0 R >>" % pages_num
    objects[pages_num - 1] = b"<< /Type /Pages /Kids [" + kids + b"] /Count %d >>" % len(page_texts)

    out = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets: list[int] = []
    for index, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % index + body + b"\nendobj\n"
    xref_pos = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for offset in offsets:
        out += b"%010d 00000 n \n" % offset
    out += b"trailer\n<< /Size %d /Root %d 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        catalog_num,
        xref_pos,
    )
    return bytes(out)


def call(method, path, body=None, token=None, raw=None, filename=None):
    """Issue one request.

    ``raw`` (bytes) plus ``filename`` sends a multipart/form-data upload;
    ``body`` sends JSON; neither sends an empty body.
    """
    url = f"{BASE}{path}"
    headers = {}
    data = None
    if raw is not None:
        boundary = "----phase2" + uuid.uuid4().hex
        head = (
            f"--{boundary}\r\n"
            f'Content-Disposition: form-data; name="file"; filename="{filename}"\r\n'
            f"Content-Type: application/octet-stream\r\n\r\n"
        ).encode("utf-8")
        data = head + raw + f"\r\n--{boundary}--\r\n".encode("utf-8")
        headers["Content-Type"] = f"multipart/form-data; boundary={boundary}"
    elif body is not None:
        data = json.dumps(body).encode()
        headers["Content-Type"] = "application/json"
    if token:
        headers["Authorization"] = f"Bearer {token}"
    req = urllib.request.Request(url, data=data, method=method, headers=headers)
    try:
        with urllib.request.urlopen(req, timeout=90) as resp:
            payload = resp.read().decode()
            return resp.status, (json.loads(payload) if payload else None)
    except urllib.error.HTTPError as e:
        payload = e.read().decode()
        try:
            return e.code, json.loads(payload)
        except ValueError:
            return e.code, payload


user = "e2e" + uuid.uuid4().hex[:8]
password = "e2epassword123"
fails: list[str] = []
uploaded_docs: list[dict] = []


def check(label, cond, extra=""):
    mark = "PASS" if cond else "FAIL"
    if not cond:
        fails.append(label)
    print(f"  [{mark}] {label}{(' -> ' + str(extra)) if extra else ''}")


def wait_task_completed(task_id, token, timeout=120):
    """Poll a task until it leaves pending/running (real-worker friendly)."""
    deadline = time.time() + timeout
    while time.time() < deadline:
        st, task = call("GET", f"/tasks/{task_id}", token=token)
        if st == 200 and task.get("status") in ("completed", "failed"):
            return task
        time.sleep(1.5)
    return {"status": "timeout"}


def upload(kb_id, filename, payload, token, label):
    st, up = call("POST", f"/knowledge-bases/{kb_id}/documents",
                  raw=payload, filename=filename, token=token)
    check(f"{label} upload 201", st == 201, f"{st} {str(up)[:160]}")
    if st != 201:
        return None
    doc, task = up["document"], up["task"]
    uploaded_docs.append(doc)
    check(f"{label} file_type", doc["file_type"] == filename.rsplit(".", 1)[-1], doc["file_type"])
    # With a real Celery broker the parse runs after the response returns;
    # wait for the terminal state before asserting on results.
    if doc["status"] != "completed":
        task = wait_task_completed(task["id"], token)
        st, docs = call("GET", f"/knowledge-bases/{kb_id}/documents", token=token)
        doc = next((d for d in docs if d["id"] == doc["id"]), doc)
        # Refresh the response payload so later checks see final results.
        up["document"], up["task"] = doc, task
    check(f"{label} completed", doc["status"] == "completed", doc["status"])
    check(f"{label} progress 100", doc["progress"] == 100, doc["progress"])
    check(f"{label} chunks > 0", doc["chunk_count"] > 0, doc["chunk_count"])
    check(f"{label} characters > 0", doc["character_count"] > 0, doc["character_count"])
    check(f"{label} task completed", task["status"] == "completed", task["status"])
    return up


print("=" * 64)
print("1) register + login")
st, reg = call("POST", "/auth/register",
               {"username": user, "email": f"{user}@example.com", "password": password})
check("register 201", st == 201, st)
st, tok = call("POST", "/auth/login", {"username": user, "password": password})
check("login 200", st == 200, st)
token = tok.get("access_token") if isinstance(tok, dict) else None
check("token issued", bool(token))

print("=" * 64)
print("2) knowledge base CRUD")
st, kbs = call("GET", "/knowledge-bases", token=token)
check("list empty initially", st == 200 and kbs == [], f"{st} {kbs}")

st, kb = call("POST", "/knowledge-bases", {"name": "产品文档", "description": "Phase 2 测试库"}, token=token)
check("create 201", st == 201, st)
kb_id = kb.get("id") if isinstance(kb, dict) else None
check("kb id returned", kb_id is not None, kb_id)
check("document_count starts 0", kb.get("document_count") == 0, kb.get("document_count"))

st, kb2 = call("PATCH", f"/knowledge-bases/{kb_id}", {"description": "已更新"}, token=token)
check("patch 200", st == 200 and kb2.get("description") == "已更新", f"{st} {kb2}")

print("=" * 64)
print("3) upload TXT (Chinese)")
txt_body = (
    "2024 年销售汇总\n\n"
    "华东区：1200 万\n"
    "华北区：980 万\n"
    "华南区：1100 万\n\n"
    "总合计 3280 万，同比增长 12%。\n"
) + ("补充说明：本季度重点关注高毛利产品线的表现。" * 40)
txt_up = upload(kb_id, "销售数据说明.txt", txt_body.encode("utf-8"), token, "txt")
if txt_up:
    check("dispatch reported", txt_up.get("dispatch") in ("inline", "queued"), txt_up.get("dispatch"))
    check("task result has chunk_count",
          isinstance(txt_up["task"].get("result"), dict)
          and "chunk_count" in (txt_up["task"].get("result") or {}),
          txt_up["task"].get("result"))

print("=" * 64)
print("4) upload DOCX (Word)")
docx_payload = build_docx([
    "本季度总销售额为 120 万元，同比增长 18%。",
    "华东区域表现最好，西南区域略有下滑。",
])
upload(kb_id, "季度报告.docx", docx_payload, token, "docx")

print("=" * 64)
print("5) upload PDF (multi-page, one empty page)")
upload(kb_id, "sales.pdf", build_pdf(["Phase two PDF page one", "", "Phase two PDF page three"]),
       token, "pdf")

if args.pdf:
    print("=" * 64)
    print(f"6) upload real-world PDF: {args.pdf}")
    with open(args.pdf, "rb") as fh:
        real_pdf = fh.read()
    upload(kb_id, "framework.pdf", real_pdf, token, "real-pdf")

expected_docs = len(uploaded_docs)

print("=" * 64)
print("7) list documents + counts")
st, docs = call("GET", f"/knowledge-bases/{kb_id}/documents", token=token)
check(f"documents listed = {expected_docs}", st == 200 and len(docs) == expected_docs,
      f"{st} n={len(docs) if isinstance(docs, list) else docs}")
st, kb3 = call("GET", f"/knowledge-bases/{kb_id}", token=token)
check(f"kb document_count = {expected_docs}", kb3.get("document_count") == expected_docs,
      kb3.get("document_count"))

print("=" * 64)
print("8) reject unsupported type")
st, bad = call("POST", f"/knowledge-bases/{kb_id}/documents",
               raw=b"MZ fake binary", filename="malware.exe", token=token)
check("exe rejected 400", st == 400, f"{st} {bad}")
st, bad = call("POST", f"/knowledge-bases/{kb_id}/documents", raw=b"", filename="empty.txt", token=token)
check("empty file rejected 400", st == 400, f"{st} {bad}")

print("=" * 64)
print("9) ownership isolation (second user cannot read)")
user2 = "e2b" + uuid.uuid4().hex[:8]
st, _ = call("POST", "/auth/register",
             {"username": user2, "email": f"{user2}@example.com", "password": password})
st, tok2 = call("POST", "/auth/login", {"username": user2, "password": password})
token2 = tok2.get("access_token")
st, _ = call("GET", f"/knowledge-bases/{kb_id}", token=token2)
check("other user gets 404", st == 404, st)
st, _ = call("GET", f"/knowledge-bases/{kb_id}/documents", token=token2)
check("other user docs 404", st == 404, st)
st, _ = call("POST", f"/knowledge-bases/{kb_id}/documents",
             raw=b"x", filename="probe.txt", token=token2)
check("other user upload 404", st == 404, st)

print("=" * 64)
print("10) unauthenticated access blocked")
st, _ = call("GET", "/knowledge-bases")
check("no token -> 401", st == 401, st)

print("=" * 64)
print("11) tasks endpoint")
st, tasks = call("GET", "/tasks", token=token)
check("tasks listed", st == 200 and len(tasks) >= expected_docs,
      f"{st} n={len(tasks) if isinstance(tasks, list) else tasks}")
if isinstance(tasks, list) and tasks:
    t = tasks[0]
    check("task has type", t.get("type") == "document.parse", t.get("type"))
    st, one = call("GET", f"/tasks/{t['id']}", token=token)
    check("single task fetch", st == 200, st)
st, filtered = call("GET", "/tasks?status=completed", token=token)
check("status filter works", st == 200 and all(t["status"] == "completed" for t in filtered), st)

print("=" * 64)
print("12) dashboard reflects real counts + recent tasks")
st, dash = call("GET", "/dashboard", token=token)
check("dashboard 200", st == 200, st)
s = dash["stats"]
check("knowledge_base_count = 1", s["knowledge_base_count"] == 1, s["knowledge_base_count"])
check(f"document_count = {expected_docs}", s["document_count"] == expected_docs, s["document_count"])
check("completed_task_count >= 1", s["completed_task_count"] >= 1, s["completed_task_count"])
breakdown = {b["date"]: b["value"] for b in dash["task_status_breakdown"]}
check("breakdown has completed>0", breakdown.get("completed", 0) >= 1, breakdown)
recent = dash.get("recent_tasks") or []
check("recent_tasks present", len(recent) >= 1, len(recent))
if recent:
    check("recent_tasks capped at 5", len(recent) <= 5, len(recent))
    check("recent task carries progress", recent[0].get("progress") == 100, recent[0].get("progress"))
    check("recent task type", recent[0].get("type") == "document.parse", recent[0].get("type"))

print("=" * 64)
print("13) cascade: delete a whole knowledge base with its document")
st, kb_cas = call("POST", "/knowledge-bases", {"name": "级联库"}, token=token)
kb_cas_id = kb_cas["id"]
st, _ = call("POST", f"/knowledge-bases/{kb_cas_id}/documents",
             raw=("级联删除回归测试。" * 30).encode("utf-8"),
             filename="cascade.txt", token=token)
check("cascade upload 201", st == 201, st)
st, _ = call("DELETE", f"/knowledge-bases/{kb_cas_id}", token=token)
check("delete kb with document 204", st == 204, st)
st, _ = call("GET", f"/knowledge-bases/{kb_cas_id}/documents", token=token)
check("documents gone with the kb (404)", st == 404, st)
st, docs_after = call("GET", f"/knowledge-bases/{kb_id}/documents", token=token)
check("unrelated kb untouched", st == 200 and len(docs_after) == expected_docs, st)

print("=" * 64)
print("14) cleanup: delete documents then the knowledge base")
for doc in uploaded_docs:
    st, _ = call("DELETE", f"/knowledge-bases/{kb_id}/documents/{doc['id']}", token=token)
    check(f"delete doc {doc['id']} 204", st == 204, st)
st, docs = call("GET", f"/knowledge-bases/{kb_id}/documents", token=token)
check("docs empty after delete", docs == [], docs)
st, _ = call("DELETE", f"/knowledge-bases/{kb_id}", token=token)
check("delete kb 204", st == 204, st)
st, kbs = call("GET", "/knowledge-bases", token=token)
check("kb list empty again", kbs == [], kbs)

print("=" * 64)
if fails:
    print(f"RESULT: {len(fails)} FAILURE(S): {fails}")
    raise SystemExit(1)
print("RESULT: ALL CHECKS PASSED")
