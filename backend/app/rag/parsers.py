"""File -> text extraction with page numbers.

Supported formats: PDF (real per-page text), DOCX and plain text (single
"page"). Every parser returns a :class:`ParsedDocument` whose pages carry
absolute character offsets so downstream chunks can be traced back to the
source file.

All failure modes surface as :class:`ValueError` with a human-readable
message — the worker stores that message on the document row and the API can
show it to the user.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

# File types this module can parse (dotless, lowercase).
SUPPORTED_FILE_TYPES: frozenset[str] = frozenset({"pdf", "docx", "txt"})

# Inserted between pages when accumulating character offsets so that the
# ranges of two consecutive pages never overlap.
PAGE_SEPARATOR: str = "\n\n"

# Decode attempts for plain text, in order. GBK is common for Chinese files
# produced on Windows; latin-1 never fails and is the last resort.
TXT_ENCODINGS: tuple[str, ...] = ("utf-8", "utf-8-sig", "gbk", "latin-1")


@dataclass
class ParsedPage:
    """One page of extracted text.

    ``page_number`` is 1-based and always mirrors the real document: pages
    that yield no text are kept with an empty ``text`` so page numbers never
    shift.
    """

    page_number: int
    text: str
    char_start: int
    char_end: int

    @property
    def length(self) -> int:
        return len(self.text)


@dataclass
class ParsedDocument:
    """Result of parsing a whole file."""

    pages: list[ParsedPage] = field(default_factory=list)
    page_count: int = 0
    character_count: int = 0
    file_type: str = ""

    def full_text(self) -> str:
        """All page texts joined with :data:`PAGE_SEPARATOR`."""
        return PAGE_SEPARATOR.join(page.text for page in self.pages)


def detect_file_type(filename: str) -> str:
    """Return the lowercase, dotless extension of ``filename``.

    Raises:
        ValueError: when the extension is missing or unsupported.
    """
    suffix = Path(str(filename)).suffix.lstrip(".").lower()
    if suffix not in SUPPORTED_FILE_TYPES:
        raise ValueError(
            f"Unsupported file type {suffix or '(none)'!r} for {filename!r}; "
            f"expected one of {', '.join(sorted(SUPPORTED_FILE_TYPES))}"
        )
    return suffix


def parse_document(path: str | Path, file_type: str) -> ParsedDocument:
    """Dispatch to the parser for ``file_type``.

    Raises:
        ValueError: unknown file type, or any error raised by the parser.
    """
    normalized = str(file_type or "").lstrip(".").lower()
    if normalized not in SUPPORTED_FILE_TYPES:
        raise ValueError(
            f"Unsupported file type {file_type!r}; "
            f"expected one of {', '.join(sorted(SUPPORTED_FILE_TYPES))}"
        )
    if normalized == "pdf":
        return parse_pdf(path)
    if normalized == "docx":
        return parse_docx(path)
    return parse_txt(path)


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
def parse_pdf(path: str | Path) -> ParsedDocument:
    """Extract per-page text from a PDF using ``pypdf``.

    Pages that yield no text are kept (with empty ``text``) so that
    ``page_number`` stays aligned with the real document.

    Raises:
        ValueError: missing file, encrypted PDF, zero pages, unreadable PDF
            or no extractable text at all.
    """
    source = _require_file(path)
    try:
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise ValueError(f"pypdf is required to parse PDF files: {exc}") from exc

    try:
        reader = PdfReader(str(source))
    except PdfReadError as exc:
        raise ValueError(f"Cannot read PDF {source.name!r}: {exc}") from exc
    except Exception as exc:  # noqa: BLE001 - surfaced as ValueError by contract
        raise ValueError(f"Cannot read PDF {source.name!r}: {exc}") from exc

    if getattr(reader, "is_encrypted", False):
        raise ValueError(f"PDF {source.name!r} is encrypted and cannot be parsed")

    try:
        raw_pages = list(reader.pages)
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"Cannot read pages of PDF {source.name!r}: {exc}") from exc

    if not raw_pages:
        raise ValueError(f"PDF {source.name!r} has no pages")

    texts: list[str] = []
    for number, page in enumerate(raw_pages, start=1):
        try:
            extracted = page.extract_text() or ""
        except Exception as exc:  # noqa: BLE001
            raise ValueError(
                f"Cannot extract text from page {number} of {source.name!r}: {exc}"
            ) from exc
        texts.append(extracted)

    if not any(text.strip() for text in texts):
        raise ValueError(
            f"PDF {source.name!r} contains no extractable text "
            "(it may be a scanned image document)"
        )

    return _build_document(texts, "pdf")


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------
def parse_docx(path: str | Path) -> ParsedDocument:
    """Extract text from a ``.docx`` file.

    DOCX has no page concept (pagination is decided by the renderer), so the
    result is always a **single page with ``page_number == 1``** containing
    every non-empty paragraph joined by newlines.

    Raises:
        ValueError: missing file, unreadable DOCX or no text at all.
    """
    source = _require_file(path)
    try:
        import docx  # type: ignore[import-untyped]
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise ValueError(f"python-docx is required to parse DOCX files: {exc}") from exc

    try:
        document = docx.Document(str(source))
        paragraphs = [para.text for para in document.paragraphs]
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"Cannot read DOCX {source.name!r}: {exc}") from exc

    body = "\n".join(text for text in paragraphs if text and text.strip())
    if not body.strip():
        raise ValueError(f"DOCX {source.name!r} contains no text")

    return _build_document([body], "docx")


# ---------------------------------------------------------------------------
# TXT
# ---------------------------------------------------------------------------
def parse_txt(path: str | Path) -> ParsedDocument:
    """Read a plain-text file.

    Encoding fallback chain: ``utf-8`` -> ``utf-8-sig`` -> ``gbk`` ->
    ``latin-1``. Same single-page model as DOCX: ``page_number == 1``.

    Raises:
        ValueError: missing file or empty content.
    """
    source = _require_file(path)
    try:
        raw_bytes = source.read_bytes()
    except OSError as exc:
        raise ValueError(f"Cannot read text file {source.name!r}: {exc}") from exc

    text, encoding = _decode_txt(raw_bytes)
    if not text.strip():
        raise ValueError(
            f"Text file {source.name!r} is empty (decoded with {encoding})"
        )

    return _build_document([text], "txt")


def _decode_txt(raw_bytes: bytes) -> tuple[str, str]:
    """Decode ``raw_bytes`` with :data:`TXT_ENCODINGS`, returning (text, encoding).

    A leading BOM is dropped so that UTF-8 and UTF-8-BOM files produce the
    same text regardless of which entry of the chain decoded them.
    """
    for encoding in TXT_ENCODINGS:
        try:
            return _strip_bom(raw_bytes.decode(encoding)), encoding
        except (UnicodeDecodeError, LookupError):
            continue
    # latin-1 is last in the chain and never raises, so this is unreachable.
    return _strip_bom(raw_bytes.decode("latin-1", errors="replace")), "latin-1"


def _strip_bom(text: str) -> str:
    return text[1:] if text.startswith("\ufeff") else text


# ---------------------------------------------------------------------------
# Shared helpers
# ---------------------------------------------------------------------------
def _require_file(path: str | Path) -> Path:
    """Resolve ``path`` and make sure it is an existing file."""
    source = Path(path)
    if not source.exists():
        raise ValueError(f"File not found: {source}")
    if not source.is_file():
        raise ValueError(f"Not a file: {source}")
    return source


def _build_document(texts: list[str], file_type: str) -> ParsedDocument:
    """Assemble a :class:`ParsedDocument` and compute absolute offsets."""
    pages: list[ParsedPage] = []
    cursor = 0
    for index, text in enumerate(texts, start=1):
        if index > 1:
            cursor += len(PAGE_SEPARATOR)
        pages.append(
            ParsedPage(
                page_number=index,
                text=text,
                char_start=cursor,
                char_end=cursor + len(text),
            )
        )
        cursor += len(text)

    return ParsedDocument(
        pages=pages,
        page_count=len(pages),
        character_count=sum(len(page.text) for page in pages),
        file_type=file_type,
    )
